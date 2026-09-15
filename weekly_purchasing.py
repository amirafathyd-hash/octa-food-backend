"""Content-driven weekly purchasing workflow and inventory cycle."""

from collections import OrderedDict
from datetime import datetime
from io import BytesIO
import json
import math
import os
import re
import secrets
import unicodedata
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

import arabic_reshaper
from flask import Blueprint, current_app, jsonify, request, send_file
import openpyxl
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
try:
    from bidi.algorithm import get_display
except ImportError:
    from bidi import get_display
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from db import execute_with_retry, get_client


weekly_purchasing_bp = Blueprint("weekly_purchasing", __name__)

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "data", "weekly_purchasing_template.xlsx")
RUN_LOG_TYPE = "weekly_purchasing_run"
INVENTORY_LOG_TYPE = "weekly_inventory_snapshot"
CONTROL_LOG_TYPE = "weekly_purchasing_control"
MAX_UPLOAD_FILES = 12
CAIRO_TZ = ZoneInfo("Africa/Cairo")
FONT_DIR = os.path.join(os.path.dirname(__file__), "fonts")
PDF_FONT = "OctaWeeklyArabic"
PDF_FONT_BOLD = "OctaWeeklyArabicBold"

HEADER_ALIASES = {
    "item": ("items", "item", "ingredient", "ingredients", "الصنف", "الاصناف", "المكون"),
    "category": ("category", "الفئة", "الفئه", "التصنيف"),
    "unit": ("unit", "الوحدة", "الوحده"),
    "daily": ("daily weight", "daily consumption", "الاستهلاك اليومي", "الوزن اليومي"),
    "weekly": ("weekly weight", "weekly consumption", "الاستهلاك الاسبوعي", "الوزن الاسبوعي"),
}


class WeeklyPurchasingError(ValueError):
    pass


def _clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.startswith("#") else re.sub(r"\s+", " ", text)


def _key(value):
    text = unicodedata.normalize("NFKC", _clean(value)).casefold()
    text = text.translate(str.maketrans("أإآىةؤئ", "ااايهوي"))
    return re.sub(r"[^a-z0-9\u0600-\u06ff]+", "", text)


def _number(value, default=0.0):
    if isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        if math.isnan(value) or math.isinf(value):
            return default
        return float(value)
    text = _clean(value).replace(",", "")
    try:
        number = float(text)
        return number if math.isfinite(number) else default
    except (TypeError, ValueError):
        return default


def _display_number(value):
    number = _number(value)
    return int(number) if number.is_integer() else round(number, 6)


def _header_role(value):
    key = _key(value)
    if not key:
        return ""
    for role, aliases in HEADER_ALIASES.items():
        if any(key == _key(alias) or _key(alias) in key for alias in aliases):
            return role
    return ""


def _find_weekly_table(ws):
    """Find the ingredient/weekly-consumption table from cell content."""
    max_row = min(ws.max_row or 0, 40)
    max_col = min(ws.max_column or 0, 40)
    if not max_row or not max_col:
        return None
    best = None
    for row_index, values in enumerate(ws.iter_rows(
        min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True,
    ), 1):
        columns = {}
        for column, value in enumerate(values, 1):
            role = _header_role(value)
            if role and role not in columns:
                columns[role] = column
        if "item" not in columns or "weekly" not in columns:
            continue
        score = 100 + 15 * len(columns)
        candidate = {"header_row": row_index, "columns": columns, "score": score}
        if best is None or candidate["score"] > best["score"]:
            best = candidate
    return best


def _read_weekly_table(ws, table):
    columns = table["columns"]
    max_col = max(columns.values())
    rows = OrderedDict()
    numeric_rows = 0
    for values in ws.iter_rows(
        min_row=table["header_row"] + 1,
        max_row=min(ws.max_row or 0, 5000),
        min_col=1,
        max_col=max_col,
        values_only=True,
    ):
        name = _clean(values[columns["item"] - 1] if len(values) >= columns["item"] else None)
        if not name or _key(name) in {"total", "totals", "الاجمالي", "المجموع"}:
            continue
        raw_weekly = values[columns["weekly"] - 1] if len(values) >= columns["weekly"] else None
        weekly = _number(raw_weekly, default=float("nan"))
        if math.isnan(weekly):
            continue
        numeric_rows += 1
        item_key = _key(name)
        if not item_key:
            continue
        category_col = columns.get("category")
        unit_col = columns.get("unit")
        category = _clean(values[category_col - 1]) if category_col and len(values) >= category_col else ""
        unit = _clean(values[unit_col - 1]) if unit_col and len(values) >= unit_col else ""
        current = rows.setdefault(item_key, {
            "key": item_key, "item": name, "category": category, "unit": unit, "weekly": 0.0,
        })
        current["weekly"] += max(0.0, weekly)
        current["category"] = current["category"] or category
        current["unit"] = current["unit"] or unit
    return rows, numeric_rows


def extract_weekly_consumption(files):
    """Aggregate every unique weekly table without using file or tab names."""
    aggregate = OrderedDict()
    sources = []
    seen_files = set()
    for upload in files:
        raw = upload.read()
        upload.seek(0)
        signature = (len(raw), raw[:4096], raw[-4096:])
        if signature in seen_files:
            continue
        seen_files.add(signature)
        try:
            workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise WeeklyPurchasingError(f"تعذر قراءة ملف Excel {upload.filename}: {exc}") from exc
        accepted_signatures = set()
        file_sources = []
        try:
            for worksheet in workbook.worksheets:
                table = _find_weekly_table(worksheet)
                if not table:
                    continue
                rows, numeric_rows = _read_weekly_table(worksheet, table)
                if numeric_rows < 2 or not rows:
                    continue
                data_signature = tuple(
                    sorted((key, round(row["weekly"], 6)) for key, row in rows.items())
                )
                if data_signature in accepted_signatures:
                    continue
                accepted_signatures.add(data_signature)
                for item_key, row in rows.items():
                    current = aggregate.setdefault(item_key, {
                        "key": item_key,
                        "item": row["item"],
                        "category": row["category"],
                        "unit": row["unit"],
                        "weekly": 0.0,
                    })
                    current["weekly"] += row["weekly"]
                    current["category"] = current["category"] or row["category"]
                    current["unit"] = current["unit"] or row["unit"]
                file_sources.append({
                    "sheet": worksheet.title,
                    "header_row": table["header_row"],
                    "items": len(rows),
                    "weekly_total": round(sum(row["weekly"] for row in rows.values()), 3),
                })
        finally:
            workbook.close()
        if file_sources:
            sources.append({"filename": upload.filename, "tables": file_sources})
    if not aggregate:
        raise WeeklyPurchasingError(
            "لم أجد جدول استهلاك أسبوعي صالح داخل الملفات. يجب أن يحتوي الجدول على عمود للصنف وعمود للاستهلاك الأسبوعي."
        )
    for row in aggregate.values():
        row["weekly"] = round(row["weekly"], 6)
    return aggregate, sources


def _catalog_headers(ws):
    best = None
    for row_index, values in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True), 1):
        keys = [_key(value) for value in values]
        if "items" not in keys:
            continue
        columns = {"item": keys.index("items") + 1}
        labels = {
            "unit": "unit",
            "category": "category",
            "base_unit": "orderbaseunitوحدهالطلبالاساسيه",
            "purchase_description": "orderunitوحدهالطلب",
            "price": "priceoforderingunitسعروحدهالطلب",
            "cost_rate": "costperkgorlالسعرللكيلواوللتر",
            "supplier": "المورد",
            "order_unit": "وحدهالطلب",
        }
        for field, wanted in labels.items():
            matching = [index + 1 for index, key in enumerate(keys) if key == wanted]
            if matching:
                columns[field] = matching[-1] if field == "order_unit" else matching[0]
        best = (row_index, columns)
        break
    return best


def load_catalog():
    if not os.path.exists(TEMPLATE_PATH):
        raise WeeklyPurchasingError("قالب Weekly Purchasing غير موجود في مجلد data")
    workbook = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        found = _catalog_headers(worksheet)
        if not found:
            raise WeeklyPurchasingError("تعذر قراءة عناوين كتالوج Weekly Purchasing")
        header_row, columns = found
        catalog = OrderedDict()
        max_col = max(columns.values())
        for values in worksheet.iter_rows(
            min_row=header_row + 1, max_row=worksheet.max_row,
            min_col=1, max_col=max_col, values_only=True,
        ):
            name = _clean(values[columns["item"] - 1])
            item_key = _key(name)
            if not item_key:
                continue
            def value(field):
                column = columns.get(field)
                return values[column - 1] if column and len(values) >= column else None
            catalog[item_key] = {
                "key": item_key,
                "item": name,
                "unit": _clean(value("unit")),
                "category": _clean(value("category")),
                "base_unit": _display_number(value("base_unit")),
                "purchase_description": _clean(value("purchase_description")),
                "price": _display_number(value("price")),
                "cost_rate": _display_number(value("cost_rate")),
                "supplier": _clean(value("supplier")),
                "order_unit": _clean(value("order_unit")),
                "is_new": False,
            }
        return catalog
    finally:
        workbook.close()


def _new_catalog_row(source):
    unit = _clean(source.get("unit"))
    unit_key = _key(unit)
    if unit_key in {"gm", "g", "جرام", "gram"}:
        base_unit, order_unit = 1000, "Kg"
    elif unit_key in {"ml", "مل", "milliliter"}:
        base_unit, order_unit = 1000, "L"
    elif unit_key in {"kg", "كيلو", "kilogram"}:
        base_unit, order_unit = 1, "Kg"
    elif unit_key in {"l", "liter", "litre", "لتر"}:
        base_unit, order_unit = 1, "L"
    else:
        base_unit, order_unit = 1, unit or "Piece"
    return {
        "key": source["key"],
        "item": source["item"],
        "unit": unit,
        "category": _clean(source.get("category")) or "صنف جديد",
        "base_unit": base_unit,
        "purchase_description": order_unit,
        "price": 0,
        "cost_rate": 0,
        "supplier": "يحتاج استكمال بيانات الشراء",
        "order_unit": order_unit,
        "is_new": True,
    }


def _log_event(file_type, file_name, payload):
    execute_with_retry(get_client().table("upload_log").insert({
        "file_type": file_type,
        "file_name": file_name,
        "item_date": payload.get("date") or None,
        "message": json.dumps(payload, ensure_ascii=False),
        "level": "info",
    }), max_attempts=2)


def _latest_payload(file_type, file_name=None):
    query = get_client().table("upload_log").select("id,file_name,message,created_at").eq("file_type", file_type)
    if file_name:
        query = query.eq("file_name", file_name)
    rows = execute_with_retry(query.order("created_at", desc=True).limit(1), max_attempts=2).data or []
    if not rows:
        return {}
    try:
        payload = json.loads(rows[0].get("message") or "{}")
    except Exception:
        return {}
    payload["_log_id"] = rows[0].get("id")
    payload["_created_at"] = rows[0].get("created_at")
    return payload


def _inventory_for_run(run_id=None, fallback_to_latest=True):
    snapshot = _latest_payload(INVENTORY_LOG_TYPE, run_id) if run_id else {}
    if snapshot or not fallback_to_latest:
        return snapshot
    return _latest_payload(INVENTORY_LOG_TYPE)


def weekly_control_for_run(run_id):
    return _latest_payload(CONTROL_LOG_TYPE, run_id)


def weekly_run_state(run_id):
    control = weekly_control_for_run(run_id)
    state = _clean(control.get("state")) or "active"
    if state == "updating":
        try:
            lock_until = datetime.fromisoformat(str(control.get("lock_until") or "").replace("Z", "+00:00"))
            if lock_until.tzinfo is None:
                lock_until = lock_until.replace(tzinfo=CAIRO_TZ)
            if datetime.now(CAIRO_TZ) >= lock_until.astimezone(CAIRO_TZ):
                state = "active"
        except (TypeError, ValueError):
            pass
    payload = _latest_payload(RUN_LOG_TYPE, run_id)
    exists = bool(payload.get("rows")) and state != "deleted"
    return {
        "exists": exists,
        "state": state if exists or state == "deleted" else "deleted",
        "revision": payload.get("revision") or payload.get("updated_at") or payload.get("created_at") or "",
        "updated_at": control.get("updated_at") or payload.get("updated_at") or payload.get("created_at") or "",
    }


def _previous_expected_stock(snapshot):
    previous_run_id = _clean(snapshot.get("run_id"))
    if not previous_run_id:
        return {}
    previous = _latest_payload(RUN_LOG_TYPE, previous_run_id)
    inventory = snapshot.get("items") or {}
    expected = {}
    for row in previous.get("rows") or []:
        key = row.get("key")
        weekly = max(0.0, _number(row.get("weekly_consumption")))
        available = max(0.0, _number(inventory.get(key)))
        order = max(0.0, weekly - available)
        expected[key] = round(available + order - weekly, 6)
    return expected


def prepare_weekly_run(files):
    consumption, sources = extract_weekly_consumption(files)
    catalog = load_catalog()
    snapshot = _inventory_for_run()
    inventory = snapshot.get("items") or {}
    previous_expected = _previous_expected_stock(snapshot)
    rows = []
    new_items = []
    ordered_keys = list(catalog)
    ordered_keys.extend(sorted((key for key in consumption if key not in catalog), key=lambda key: consumption[key]["item"].casefold()))
    for item_key in ordered_keys:
        source = consumption.get(item_key) or {}
        master = dict(catalog.get(item_key) or _new_catalog_row(source))
        if master.get("is_new"):
            new_items.append(master["item"])
        master["category"] = master.get("category") or source.get("category") or ""
        master["unit"] = master.get("unit") or source.get("unit") or ""
        master["weekly_consumption"] = _display_number(source.get("weekly"))
        master["expected_stock"] = _display_number(previous_expected.get(item_key)) if item_key in previous_expected else ""
        master["available_stock"] = _display_number(inventory.get(item_key)) if item_key in inventory else ""
        rows.append(master)
    now = datetime.now(CAIRO_TZ)
    run_id = secrets.token_urlsafe(14)
    payload = {
        "run_id": run_id,
        "date": now.date().isoformat(),
        "created_at": now.isoformat(),
        "rows": rows,
        "sources": sources,
        "new_items": new_items,
        "inventory_source_run_id": snapshot.get("run_id") or "",
        "inventory_updated_at": snapshot.get("submitted_at") or snapshot.get("created_at") or "",
        "revision": 1,
    }
    _log_event(RUN_LOG_TYPE, run_id, payload)
    return payload


def _safe_sheet_title(value):
    text = re.sub(r"[\\/*?:\[\]]+", " ", _clean(value))[:31].strip()
    return text or "Weekly Purchasing"


def build_weekly_workbook(run_payload, inventory_override=None):
    rows = run_payload.get("rows") or []
    if not rows:
        raise WeeklyPurchasingError("لا توجد أصناف لإنشاء Weekly Purchasing")
    inventory_override = inventory_override or {}
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = _safe_sheet_title("Weekly Purchasing")
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.rightToLeft = False
    worksheet.sheet_view.zoomScale = 75
    worksheet.sheet_properties.tabColor = "0F6765"
    worksheet.freeze_panes = "D9"

    navy = "17324D"
    teal = "0F6765"
    orange = "F4A340"
    cream = "FFF9F0"
    pale_teal = "EAF5F2"
    pale_orange = "FFF1DA"
    pale_red = "FDE9E7"
    green = "2F7D4A"
    grey = "667784"
    white = "FFFFFF"
    light_line = "D7E2DF"
    thin = Side(style="thin", color=light_line)
    bottom = Side(style="medium", color=teal)

    last_row = 8 + len(rows)
    for column in range(1, 18):
        worksheet.cell(1, column).fill = PatternFill("solid", fgColor=orange if column <= 2 else navy)
    worksheet.row_dimensions[1].height = 8
    for row_number in (2, 3, 4):
        worksheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=17)
        for column in range(1, 18):
            worksheet.cell(row_number, column).fill = PatternFill("solid", fgColor=navy)
    worksheet["A2"] = "WEEKLY PURCHASING  |  المشتريات الأسبوعية"
    worksheet["A2"].font = Font(name="Arial", size=22, bold=True, color=white)
    worksheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet["A3"] = "خطة شراء ذكية محدثة من استهلاك محطات التشغيل والمخزون الفعلي"
    worksheet["A3"].font = Font(name="Arial", size=11, color="CFE7E3")
    worksheet["A3"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet["A4"] = f"تاريخ التجهيز: {run_payload.get('date') or ''}   •   أدخل المخزون في العمود الأصفر وستتحدث الكميات والتكلفة تلقائيًا"
    worksheet["A4"].font = Font(name="Arial", size=9, bold=True, color="FFD99F")
    worksheet["A4"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[2].height = 36
    worksheet.row_dimensions[3].height = 22
    worksheet.row_dimensions[4].height = 25

    cards = (
        (1, 4, "عدد الأصناف", f"=COUNTA(A9:A{last_row})", "0"),
        (5, 8, "أصناف مطلوب شراؤها", f'=COUNTIF(L9:L{last_row},">0")', "0"),
        (9, 12, "إجمالي تكلفة الطلب", f"=SUM(Q9:Q{last_row})", '#,##0.00 "ر.س"'),
        (13, 17, "آخر تحديث للمخزون", run_payload.get("inventory_updated_at") or "لم يسجل بعد", "General"),
    )
    for start_column, end_column, label, value, number_format in cards:
        worksheet.merge_cells(start_row=5, start_column=start_column, end_row=5, end_column=end_column)
        worksheet.merge_cells(start_row=6, start_column=start_column, end_row=6, end_column=end_column)
        label_cell = worksheet.cell(5, start_column, label)
        value_cell = worksheet.cell(6, start_column, value)
        for row_number in (5, 6):
            for column in range(start_column, end_column + 1):
                cell = worksheet.cell(row_number, column)
                cell.fill = PatternFill("solid", fgColor=pale_teal)
                cell.border = Border(bottom=bottom)
        label_cell.font = Font(name="Arial", size=9, bold=True, color=teal)
        value_cell.font = Font(name="Arial", size=14, bold=True, color=navy)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.number_format = number_format
    worksheet.row_dimensions[5].height = 22
    worksheet.row_dimensions[6].height = 31

    worksheet.merge_cells("A7:Q7")
    worksheet["A7"] = "الخلايا الصفراء فقط لإدخال المخزون. الأصناف الجديدة تظهر تلقائيًا باللون الأحمر حتى استكمال المورد والسعر."
    worksheet["A7"].fill = PatternFill("solid", fgColor=pale_orange)
    worksheet["A7"].font = Font(name="Arial", size=9, bold=True, color=navy)
    worksheet["A7"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[7].height = 24

    headers = [
        "ITEMS", "Unit", "Category", "Order Base Unit\nوحدة الطلب الأساسية",
        "Order Unit\nوحدة الطلب", "Price of ordering unit\nسعر وحدة الطلب",
        "Cost Per KG or L\nالسعر للكيلو أو اللتر", "الاستهلاك الأسبوعي",
        "الحد الأدنى للكمية المتاحة (MAQ)", "المخزون المتوقع المتاح",
        "المخزون المتاح", "الطلب الأسبوعي", "المخزون المتوقع للأسبوع القادم",
        "المورد", "طلب الأسبوع", "وحدة الطلب", "سعر الطلب الأسبوعي",
    ]
    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(8, column, header)
        cell.fill = PatternFill("solid", fgColor=teal if column >= 8 else navy)
        if column == 11:
            cell.fill = PatternFill("solid", fgColor=orange)
        elif column in (12, 15, 16, 17):
            cell.fill = PatternFill("solid", fgColor=green)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=Side(style="thin", color=white), bottom=bottom)
    worksheet.row_dimensions[8].height = 64

    for index, row in enumerate(rows, 9):
        item_key = row.get("key") or _key(row.get("item"))
        available = inventory_override.get(item_key, row.get("available_stock", ""))
        values = [
            row.get("item"), row.get("unit"), row.get("category"), row.get("base_unit"),
            row.get("purchase_description"), row.get("price"), row.get("cost_rate"),
            row.get("weekly_consumption"), f"=H{index}", row.get("expected_stock", ""),
            _display_number(available) if available != "" else "",
            f'=MAX(0,H{index}-K{index})',
            f'=MAX(0,K{index}+L{index}-H{index})',
            row.get("supplier"), f'=IF(D{index}=0,0,L{index}/D{index})', row.get("order_unit"),
            f'=O{index}*F{index}',
        ]
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(index, column, value)
            cell.font = Font(name="Arial", size=11, color=navy, bold=True)
            cell.alignment = Alignment(
                horizontal="left" if column in (1, 3, 5, 14, 16) else "right" if column in (4, 6, 7, 8, 9, 10, 11, 12, 13, 15, 17) else "center",
                vertical="center",
                wrap_text=True,
            )
            cell.fill = PatternFill("solid", fgColor=cream if index % 2 else white)
            if column == 11:
                cell.fill = PatternFill("solid", fgColor=pale_orange)
                cell.protection = Protection(locked=False)
            elif row.get("is_new") and column in (1, 4, 5, 6, 7, 14, 16):
                cell.fill = PatternFill("solid", fgColor=pale_red)
            elif column in (12, 15, 17):
                cell.fill = PatternFill("solid", fgColor=pale_teal)
            cell.border = Border(bottom=thin)
            if column in (4, 6, 7, 8, 9, 10, 11, 12, 13):
                cell.number_format = '#,##0.000'
            elif column in (15, 17):
                cell.number_format = '#,##0.00'
        worksheet.row_dimensions[index].height = 38

    worksheet.auto_filter.ref = f"A8:Q{last_row}"
    worksheet.conditional_formatting.add(
        f"K9:K{last_row}",
        FormulaRule(formula=["K9=\"\""], fill=PatternFill("solid", fgColor=pale_orange)),
    )
    worksheet.conditional_formatting.add(
        f"L9:L{last_row}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="DDF1E5")),
    )
    worksheet.conditional_formatting.add(
        f"N9:N{last_row}",
        FormulaRule(formula=['ISNUMBER(SEARCH("يحتاج استكمال",N9))'], fill=PatternFill("solid", fgColor=pale_red)),
    )

    widths = (38, 11, 18, 18, 24, 18, 18, 18, 19, 19, 18, 18, 22, 30, 18, 20, 20)
    for column, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "8:8"
    worksheet.print_area = f"A1:Q{last_row}"
    worksheet.protection.sheet = False
    workbook.calculation.calcMode = "auto"
    workbook.calculation.calcId = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook


def _register_pdf_fonts():
    regular = os.path.join(FONT_DIR, "IBMPlexSansArabic-Regular.ttf")
    bold = os.path.join(FONT_DIR, "IBMPlexSansArabic-Bold.ttf")
    if not os.path.isfile(regular) or not os.path.isfile(bold):
        raise WeeklyPurchasingError("خط PDF العربي غير موجود على الخادم")
    if PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT, regular))
    if PDF_FONT_BOLD not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT_BOLD, bold))


def _pdf_text(value):
    text = _clean(value)
    if re.search(r"[\u0600-\u06ff]", text):
        text = get_display(arabic_reshaper.reshape(text))
    return escape(text)


def build_weekly_pdf(run_payload, inventory):
    """Build a complete, multi-page purchasing PDF after inventory submission."""
    _register_pdf_fonts()
    rows = run_payload.get("rows") or []
    if not rows:
        raise WeeklyPurchasingError("لا توجد بيانات لإنشاء PDF")
    output = BytesIO()
    page_size = landscape(A3)
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=13 * mm,
        title="Weekly Purchasing",
        author="Octa Food",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "WeeklyTitle", parent=styles["Title"], fontName=PDF_FONT_BOLD,
        fontSize=24, leading=30, textColor=colors.HexColor("#17324D"), alignment=TA_CENTER,
    )
    subtitle_style = ParagraphStyle(
        "WeeklySubtitle", parent=styles["Normal"], fontName=PDF_FONT,
        fontSize=11, leading=16, textColor=colors.HexColor("#536872"), alignment=TA_CENTER,
    )
    cell_style = ParagraphStyle(
        "WeeklyCell", parent=styles["Normal"], fontName=PDF_FONT_BOLD,
        fontSize=8.5, leading=11, textColor=colors.HexColor("#17324D"), alignment=TA_RIGHT,
    )
    center_style = ParagraphStyle(
        "WeeklyCenter", parent=cell_style, alignment=TA_CENTER,
    )
    header_style = ParagraphStyle(
        "WeeklyHeader", parent=center_style, fontName=PDF_FONT_BOLD,
        fontSize=9, leading=11, textColor=colors.white,
    )

    story = [
        Paragraph(_pdf_text("المشتريات الأسبوعية - WEEKLY PURCHASING"), title_style),
        Paragraph(_pdf_text(
            f"تاريخ الدورة: {run_payload.get('date') or ''} | عدد الأصناف: {len(rows)} | محسوب بعد إدخال المخزون الفعلي"
        ), subtitle_style),
        Spacer(1, 7 * mm),
    ]
    headers = ["الصنف", "التصنيف", "الاستهلاك", "المخزون", "الطلب", "المتوقع", "المورد", "وحدة الطلب", "تكلفة الطلب"]
    data = [[Paragraph(_pdf_text(value), header_style) for value in headers]]
    total_cost = 0.0
    for row in rows:
        key = row.get("key") or _key(row.get("item"))
        weekly = max(0.0, _number(row.get("weekly_consumption")))
        available = max(0.0, _number(inventory.get(key)))
        order = max(0.0, weekly - available)
        expected = max(0.0, available + order - weekly)
        base_unit = _number(row.get("base_unit"))
        order_units = order / base_unit if base_unit > 0 else 0.0
        cost = order_units * max(0.0, _number(row.get("price")))
        total_cost += cost
        values = [
            row.get("item"), row.get("category"), f"{weekly:,.3f}", f"{available:,.3f}",
            f"{order:,.3f}", f"{expected:,.3f}", row.get("supplier"),
            f"{order_units:,.2f} {row.get('order_unit') or ''}", f"{cost:,.2f}",
        ]
        data.append([
            Paragraph(_pdf_text(value), cell_style if index in (0, 1, 6, 7) else center_style)
            for index, value in enumerate(values)
        ])
    data.append([
        Paragraph(_pdf_text("الإجمالي"), header_style), "", "", "", "", "", "", "",
        Paragraph(f"{total_cost:,.2f}", header_style),
    ])
    table = Table(
        data,
        repeatRows=1,
        colWidths=[70 * mm, 29 * mm, 27 * mm, 27 * mm, 27 * mm, 27 * mm, 48 * mm, 34 * mm, 30 * mm],
        hAlign="CENTER",
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F6765")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#17324D")),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("SPAN", (0, -1), (7, -1)),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D9D5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F3F8F6")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(PDF_FONT, 8)
        canvas.setFillColor(colors.HexColor("#687B84"))
        canvas.drawCentredString(page_size[0] / 2, 6 * mm, f"Octa Food | Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return output


def _run_payload(run_id):
    payload = _latest_payload(RUN_LOG_TYPE, run_id)
    if not payload.get("rows"):
        raise WeeklyPurchasingError("رابط Weekly Purchasing غير موجود أو انتهت بياناته")
    return payload


def workbook_for_run(run_id):
    payload = _run_payload(run_id)
    # A historical download must only use inventory submitted for that run.
    # Otherwise, retain the stock values frozen into the run when it was created.
    snapshot = _inventory_for_run(run_id, fallback_to_latest=False)
    inventory = snapshot.get("items") or {}
    workbook = build_weekly_workbook(payload, inventory_override=inventory)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output, payload


def weekly_purchasing_download(files):
    payload = prepare_weekly_run(files)
    output, _ = workbook_for_run(payload["run_id"])
    return output, f"Weekly_Purchasing_{payload['date']}.xlsx"


@weekly_purchasing_bp.route("/api/weekly-purchasing/prepare", methods=["POST"])
def weekly_purchasing_prepare():
    files = [file for file in request.files.getlist("files") if file and file.filename]
    if not 1 <= len(files) <= MAX_UPLOAD_FILES:
        return jsonify({"error": f"ارفع من ملف واحد إلى {MAX_UPLOAD_FILES} ملف Excel"}), 400
    if any(not file.filename.lower().endswith((".xlsx", ".xlsm")) for file in files):
        return jsonify({"error": "مسموح فقط بملفات XLSX وXLSM"}), 400
    try:
        payload = prepare_weekly_run(files)
        run_id = payload["run_id"]
        return jsonify({
            "ok": True,
            "run_id": run_id,
            "items_count": len(payload["rows"]),
            "new_items_count": len(payload["new_items"]),
            "new_items": payload["new_items"],
            "source_tables_count": sum(len(source["tables"]) for source in payload["sources"]),
            "download_path": f"/api/weekly-purchasing/{run_id}/xlsx",
            "inventory_path": f"/weekly-inventory.html?id={run_id}&v=27",
        })
    except WeeklyPurchasingError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        current_app.logger.exception("weekly purchasing prepare failed")
        return jsonify({"error": f"تعذر تجهيز Weekly Purchasing: {str(exc)[:180]}"}), 500


@weekly_purchasing_bp.route("/api/weekly-purchasing/<run_id>/xlsx", methods=["GET"])
def weekly_purchasing_xlsx(run_id):
    try:
        output, payload = workbook_for_run(run_id)
        return send_file(
            output,
            as_attachment=True,
            download_name=f"Weekly_Purchasing_{payload.get('date') or 'week'}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except WeeklyPurchasingError as exc:
        return jsonify({"error": str(exc)}), 404
    except Exception as exc:
        current_app.logger.exception("weekly purchasing download failed")
        return jsonify({"error": f"تعذر إنشاء ملف Excel: {str(exc)[:180]}"}), 500


@weekly_purchasing_bp.route("/api/weekly-purchasing/<run_id>/pdf", methods=["GET"])
def weekly_purchasing_pdf(run_id):
    try:
        payload = _run_payload(run_id)
        snapshot = _inventory_for_run(run_id, fallback_to_latest=False)
        inventory = snapshot.get("items") or {}
        if len(inventory) != len(payload.get("rows") or []):
            return jsonify({"error": "ملف PDF يصبح متاحًا بعد اكتمال إدخال المخزون"}), 409
        output = build_weekly_pdf(payload, inventory)
        return send_file(
            output,
            as_attachment=True,
            download_name=f"Weekly_Purchasing_{payload.get('date') or 'week'}.pdf",
            mimetype="application/pdf",
        )
    except WeeklyPurchasingError as exc:
        return jsonify({"error": str(exc)}), 404
    except Exception as exc:
        current_app.logger.exception("weekly purchasing PDF failed")
        return jsonify({"error": f"تعذر إنشاء ملف PDF: {str(exc)[:180]}"}), 500


@weekly_purchasing_bp.route("/api/weekly-inventory/<run_id>", methods=["GET"])
def weekly_inventory_get(run_id):
    try:
        state = weekly_run_state(run_id)
        if not state["exists"] or state["state"] != "active":
            return jsonify({"error": "جاري تحديث البيانات", **state}), 409
        payload = _run_payload(run_id)
        run_snapshot = _inventory_for_run(run_id, fallback_to_latest=False)
        snapshot = run_snapshot or _inventory_for_run()
        inventory = snapshot.get("items") or {}
        row_count = len(payload.get("rows") or [])
        completed = bool(run_snapshot.get("submitted_at")) and len(run_snapshot.get("items") or {}) == row_count
        return jsonify({
            "ok": True,
            "run_id": run_id,
            "date": payload.get("date"),
            "created_at": payload.get("created_at"),
            "last_submitted_at": snapshot.get("submitted_at") or "",
            "last_worker_name": snapshot.get("worker_name") or "",
            "completed": completed,
            "download_path": f"/api/weekly-purchasing/{run_id}/xlsx",
            "pdf_path": f"/api/weekly-purchasing/{run_id}/pdf" if completed else "",
            "revision": state["revision"],
            "rows": [{
                "key": row.get("key"),
                "item": row.get("item"),
                "category": row.get("category"),
                "unit": row.get("unit"),
                "weekly_consumption": row.get("weekly_consumption"),
                "available_stock": inventory.get(row.get("key"), row.get("available_stock", "")),
            } for row in payload.get("rows") or []],
        })
    except WeeklyPurchasingError as exc:
        return jsonify({"error": str(exc)}), 404
    except Exception as exc:
        current_app.logger.exception("weekly inventory load failed")
        return jsonify({"error": f"تعذر تحميل رابط المخزون: {str(exc)[:180]}"}), 500


@weekly_purchasing_bp.route("/api/weekly-inventory/<run_id>", methods=["POST"])
def weekly_inventory_submit(run_id):
    try:
        state = weekly_run_state(run_id)
        if not state["exists"] or state["state"] != "active":
            return jsonify({"error": "جاري تحديث البيانات", **state}), 409
        run_payload = _run_payload(run_id)
        incoming = request.get_json(silent=True) or {}
        worker_name = _clean(incoming.get("worker_name"))
        if not worker_name:
            return jsonify({"error": "اكتب اسم العامل المسؤول عن الجرد"}), 400
        allowed = {row.get("key") for row in run_payload.get("rows") or []}
        items = {}
        for row in incoming.get("rows") or []:
            item_key = _clean(row.get("key"))
            if item_key not in allowed:
                continue
            raw_value = row.get("available_stock")
            if raw_value in (None, ""):
                continue
            number = _number(raw_value, default=float("nan"))
            if math.isnan(number) or number < 0:
                return jsonify({"error": f"كمية المخزون غير صحيحة للصنف {row.get('item') or item_key}"}), 400
            items[item_key] = _display_number(number)
        if len(items) != len(allowed):
            return jsonify({
                "error": f"أكمل مخزون كل الأصناف أولًا ({len(items)} من {len(allowed)})"
            }), 400
        now = datetime.now(CAIRO_TZ).isoformat()
        payload = {
            "run_id": run_id,
            "date": run_payload.get("date"),
            "created_at": now,
            "submitted_at": now,
            "worker_name": worker_name,
            "items_count": len(items),
            "items": items,
        }
        _log_event(INVENTORY_LOG_TYPE, run_id, payload)
        try:
            from appointments_api import send_push_to_all
            send_push_to_all(
                "تم اكتمال جرد المخزون الأسبوعي",
                f"{worker_name} سجّل مخزون {len(items)} صنف",
                f"weekly-receiving?run_id={run_id}",
            )
        except Exception:
            current_app.logger.warning("weekly inventory push notification failed", exc_info=True)
        return jsonify({
            "ok": True,
            "submitted_at": now,
            "items_count": len(items),
            "download_path": f"/api/weekly-purchasing/{run_id}/xlsx",
            "pdf_path": f"/api/weekly-purchasing/{run_id}/pdf",
        })
    except WeeklyPurchasingError as exc:
        return jsonify({"error": str(exc)}), 404
    except Exception as exc:
        current_app.logger.exception("weekly inventory submit failed")
        return jsonify({"error": f"تعذر حفظ المخزون: {str(exc)[:180]}"}), 500


@weekly_purchasing_bp.route("/api/weekly-inventory/<run_id>/status", methods=["GET"])
def weekly_inventory_status(run_id):
    state = weekly_run_state(run_id)
    return jsonify(state), 200 if state["exists"] and state["state"] == "active" else 409
